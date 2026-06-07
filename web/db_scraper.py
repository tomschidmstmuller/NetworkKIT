#!/usr/bin/env python3
"""
Advanced Database Web Scraper
Extracts data from websites and exports to Excel, MongoDB, or SQL databases
"""

import requests
from bs4 import BeautifulSoup
import time
import json
import csv
import sqlite3
import pymongo
from urllib.parse import urljoin, urlparse
import re
from datetime import datetime
import argparse
import logging
import hashlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import sys
from typing import Dict, List, Any, Optional
import os

class DatabaseWebScraper:
    def __init__(self, config: Dict = None):
        """Initialize scraper with configuration"""
        self.config = config or {}
        self.session = requests.Session()
        self.session.headers.update({
            'User-Agent': 'DatabaseWebScraper/1.0 (Educational Purpose)'
        })
        self.data = []
        self.metadata = {
            'scraped_at': datetime.now().isoformat(),
            'total_records': 0,
            'source_urls': []
        }
        
        # Setup logging
        logging.basicConfig(
            level=logging.INFO,
            format='%(asctime)s - %(levelname)s - %(message)s',
            handlers=[
                logging.FileHandler('scraper.log'),
                logging.StreamHandler()
            ]
        )
        self.logger = logging.getLogger(__name__)
        
        self.banner = """
   ▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄▄
  █▄░▄█░▄▄░█▄░▄█░▄▄░█░▄▄▀█▄░▄█
  █░▀░█▄▄▄░██░██▄▄▄░█░▀▀▄██░██
  ▀░░░▀▀▀▀▀▀░░░▀▀▀▀▀▀▀▀▀▀░░░▀▀
  ██▄ ███ ▄▄▄ █▀▄ ▄▀█ ▄▄█ ▄▄▀█
  █▄█ ██▄ ▀▀▄ █░▀▀░█ ▄▄█ ▀▀▄██
  ▀▀▀ ▀▀▀ ▀▀▀ ▀░░░▀▀▀▀▀▀▀▀▀▀░░
  UNIVERSAL SPACE MARINE INTELLIGENT
        """
    
    def scrape_website(self, url: str, selectors: Dict[str, str], pagination: Dict = None):
        """
        Scrape website using CSS selectors
        
        Args:
            url: Target URL
            selectors: Dictionary mapping field names to CSS selectors
            pagination: Pagination settings {'next_button': 'css_selector', 'max_pages': int}
        """
        self.logger.info(f"Starting scrape of {url}")
        self.metadata['source_urls'].append(url)
        
        current_url = url
        page = 1
        
        while current_url and page <= pagination.get('max_pages', 1):
            self.logger.info(f"Scraping page {page}: {current_url}")
            
            response = self.get_page(current_url)
            if not response:
                break
            
            soup = BeautifulSoup(response.content, 'html.parser')
            
            # Extract data using selectors
            items = self.extract_items(soup, selectors)
            
            if items:
                for item in items:
                    item['_source_url'] = current_url
                    item['_scraped_at'] = datetime.now().isoformat()
                    self.data.append(item)
                
                self.logger.info(f"Extracted {len(items)} items from page {page}")
            else:
                self.logger.warning(f"No items found on page {page}")
            
            # Handle pagination
            if pagination and page < pagination.get('max_pages', 1):
                next_url = self.get_next_page_url(soup, pagination.get('next_button'))
                if next_url:
                    current_url = urljoin(current_url, next_url)
                    page += 1
                    time.sleep(pagination.get('delay', 1))
                else:
                    break
            else:
                break
        
        self.metadata['total_records'] = len(self.data)
        self.logger.info(f"Scraping complete. Total records: {len(self.data)}")
        return self.data
    
    def scrape_table(self, url: str, table_index: int = 0):
        """Scrape HTML table from website"""
        self.logger.info(f"Scraping table from {url}")
        
        response = self.get_page(url)
        if not response:
            return []
        
        soup = BeautifulSoup(response.content, 'html.parser')
        tables = soup.find_all('table')
        
        if table_index >= len(tables):
            self.logger.error(f"Table index {table_index} not found. Only {len(tables)} tables exist.")
            return []
        
        table = tables[table_index]
        
        # Extract headers
        headers = []
        header_row = table.find('tr')
        if header_row:
            headers = [th.get_text(strip=True) for th in header_row.find_all(['th', 'td'])]
        
        # Extract data rows
        rows = []
        for tr in table.find_all('tr')[1:]:  # Skip header row
            row_data = {}
            cells = tr.find_all(['td', 'th'])
            
            for i, cell in enumerate(cells):
                if i < len(headers):
                    row_data[headers[i]] = cell.get_text(strip=True)
                else:
                    row_data[f'column_{i}'] = cell.get_text(strip=True)
            
            if row_data:
                rows.append(row_data)
        
        self.data = rows
        self.metadata['total_records'] = len(rows)
        self.logger.info(f"Extracted {len(rows)} rows from table")
        return rows
    
    def scrape_api_endpoint(self, api_url: str, params: Dict = None, auth: tuple = None):
        """Scrape data from API endpoint"""
        self.logger.info(f"Scraping API: {api_url}")
        
        try:
            response = self.session.get(api_url, params=params, auth=auth, timeout=10)
            response.raise_for_status()
            
            # Try to parse JSON
            data = response.json()
            
            # Handle different JSON structures
            if isinstance(data, list):
                self.data = data
            elif isinstance(data, dict):
                if 'data' in data:
                    self.data = data['data']
                elif 'results' in data:
                    self.data = data['results']
                else:
                    self.data = [data]
            
            self.metadata['total_records'] = len(self.data)
            self.logger.info(f"Extracted {len(self.data)} records from API")
            return self.data
            
        except Exception as e:
            self.logger.error(f"API scraping failed: {e}")
            return []
    
    def extract_items(self, soup: BeautifulSoup, selectors: Dict[str, str]) -> List[Dict]:
        """Extract items using CSS selectors"""
        items = []
        
        # Find all item containers
        container_selector = selectors.get('_container', 'body')
        containers = soup.select(container_selector)
        
        for container in containers:
            item = {}
            
            for field, selector in selectors.items():
                if field.startswith('_'):
                    continue  # Skip special fields
                
                try:
                    element = container.select_one(selector)
                    if element:
                        # Handle different data types
                        if selector.endswith('::text'):
                            value = element.get_text(strip=True)
                        elif selector.endswith('::attr(href)'):
                            value = element.get('href', '')
                        elif selector.endswith('::attr(src)'):
                            value = element.get('src', '')
                        else:
                            value = element.get_text(strip=True)
                        
                        # Clean and format
                        value = self.clean_value(value)
                        item[field] = value
                except Exception as e:
                    self.logger.debug(f"Error extracting {field}: {e}")
                    item[field] = None
            
            if item:
                items.append(item)
        
        return items
    
    def get_page(self, url: str, retries: int = 3) -> Optional[requests.Response]:
        """Fetch webpage with retries"""
        for attempt in range(retries):
            try:
                response = self.session.get(url, timeout=10)
                response.raise_for_status()
                return response
            except requests.RequestException as e:
                self.logger.warning(f"Attempt {attempt + 1} failed: {e}")
                time.sleep(2)
        
        return None
    
    def get_next_page_url(self, soup: BeautifulSoup, next_selector: str) -> Optional[str]:
        """Find next page URL"""
        if not next_selector:
            return None
        
        next_button = soup.select_one(next_selector)
        if next_button:
            return next_button.get('href')
        
        return None
    
    def clean_value(self, value: str) -> str:
        """Clean extracted value"""
        if not value:
            return None
        
        # Remove extra whitespace
        value = re.sub(r'\s+', ' ', value).strip()
        
        # Remove special characters (optional)
        # value = re.sub(r'[^\w\s\-.,@:()]', '', value)
        
        return value
    
    # DATABASE EXPORT METHODS
    
    def export_to_excel(self, filename: str = None, sheet_name: str = 'Data'):
        """Export data to Excel file"""
        if not self.data:
            self.logger.warning("No data to export")
            return False
        
        try:
            # Convert to DataFrame
            df = pd.DataFrame(self.data)
            
            # Generate filename if not provided
            if not filename:
                filename = f"export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            # Create Excel writer with openpyxl engine
            with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                
                # Get workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets[sheet_name]
                
                # Style the header
                header_font = Font(bold=True, color='FFFFFF')
                header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
                header_alignment = Alignment(horizontal='center', vertical='center')
                
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
            
            # Add metadata sheet
            with pd.ExcelWriter(filename, engine='openpyxl', mode='a') as writer:
                metadata_df = pd.DataFrame([self.metadata])
                metadata_df.to_excel(writer, sheet_name='Metadata', index=False)
            
            self.logger.info(f"Data exported to Excel: {filename}")
            return True
            
        except Exception as e:
            self.logger.error(f"Excel export failed: {e}")
            return False
    
    def export_to_mongodb(self, connection_string: str, database_name: str, collection_name: str):
        """Export data to MongoDB"""
        try:
            # Connect to MongoDB
            client = pymongo.MongoClient(connection_string)
            db = client[database_name]
            collection = db[collection_name]
            
            # Insert data
            if self.data:
                # Add metadata to each document
                for item in self.data:
                    item['_exported_at'] = datetime.now()
                
                result = collection.insert_many(self.data)
                self.logger.info(f"Inserted {len(result.inserted_ids)} documents into MongoDB")
                
                # Create indexes for better query performance
                collection.create_index([('_scraped_at', pymongo.DESCENDING)])
                
                # Store metadata
                db['scraper_metadata'].insert_one(self.metadata)
                
            client.close()
            return True
            
        except Exception as e:
            self.logger.error(f"MongoDB export failed: {e}")
            return False
    
    def export_to_sqlite(self, filename: str = 'scraped_data.db', table_name: str = 'scraped_data'):
        """Export data to SQLite database"""
        try:
            conn = sqlite3.connect(filename)
            cursor = conn.cursor()
            
            if not self.data:
                self.logger.warning("No data to export")
                return False
            
            # Analyze data structure
            sample = self.data[0]
            columns = list(sample.keys())
            
            # Create table
            column_defs = []
            for col in columns:
                col_type = self.infer_sqlite_type(sample[col])
                column_defs.append(f'"{col}" {col_type}')
            
            create_table_sql = f'''
                CREATE TABLE IF NOT EXISTS "{table_name}" (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    {', '.join(column_defs)}
                )
            '''
            
            cursor.execute(create_table_sql)
            
            # Insert data
            for item in self.data:
                placeholders = ', '.join(['?' for _ in columns])
                values = [item.get(col) for col in columns]
                
                insert_sql = f'INSERT INTO "{table_name}" ({", ".join([f"\"{col}\"" for col in columns])}) VALUES ({placeholders})'
                cursor.execute(insert_sql, values)
            
            # Create metadata table
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS scraper_metadata (
                    key TEXT PRIMARY KEY,
                    value TEXT,
                    exported_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            ''')
            
            for key, value in self.metadata.items():
                cursor.execute('INSERT OR REPLACE INTO scraper_metadata (key, value) VALUES (?, ?)', 
                             (key, json.dumps(value) if not isinstance(value, str) else value))
            
            conn.commit()
            conn.close()
            
            self.logger.info(f"Data exported to SQLite: {filename} (table: {table_name})")
            return True
            
        except Exception as e:
            self.logger.error(f"SQLite export failed: {e}")
            return False
    
    def export_to_postgresql(self, host: str, database: str, user: str, password: str, 
                            table_name: str, port: int = 5432):
        """Export data to PostgreSQL"""
        try:
            import psycopg2
            from psycopg2.extras import execute_values
            
            conn = psycopg2.connect(
                host=host,
                database=database,
                user=user,
                password=password,
                port=port
            )
            cursor = conn.cursor()
            
            if not self.data:
                self.logger.warning("No data to export")
                return False
            
            # Analyze data structure
            columns = list(self.data[0].keys())
            
            # Create table
            column_defs = []
            for col in columns:
                col_type = self.infer_postgres_type(self.data[0][col])
                column_defs.append(f'"{col}" {col_type}')
            
            create_table_sql = f'''
                CREATE TABLE IF NOT EXISTS "{table_name}" (
                    id SERIAL PRIMARY KEY,
                    {', '.join(column_defs)},
                    exported_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            '''
            
            cursor.execute(create_table_sql)
            
            # Insert data using execute_values for efficiency
            values = [tuple(item.get(col) for col in columns) for item in self.data]
            insert_sql = f'INSERT INTO "{table_name}" ({", ".join([f"\"{col}\"" for col in columns])}) VALUES %s'
            execute_values(cursor, insert_sql, values)
            
            conn.commit()
            cursor.close()
            conn.close()
            
            self.logger.info(f"Data exported to PostgreSQL: {database}.{table_name}")
            return True
            
        except ImportError:
            self.logger.error("psycopg2 not installed. Run: pip install psycopg2-binary")
            return False
        except Exception as e:
            self.logger.error(f"PostgreSQL export failed: {e}")
            return False
    
    def export_to_mysql(self, host: str, database: str, user: str, password: str, 
                       table_name: str, port: int = 3306):
        """Export data to MySQL/MariaDB"""
        try:
            import mysql.connector
            from mysql.connector import Error
            
            conn = mysql.connector.connect(
                host=host,
                database=database,
                user=user,
                password=password,
                port=port
            )
            cursor = conn.cursor()
            
            if not self.data:
                self.logger.warning("No data to export")
                return False
            
            # Analyze data structure
            columns = list(self.data[0].keys())
            
            # Create table
            column_defs = []
            for col in columns:
                col_type = self.infer_mysql_type(self.data[0][col])
                column_defs.append(f'`{col}` {col_type}')
            
            create_table_sql = f'''
                CREATE TABLE IF NOT EXISTS `{table_name}` (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    {', '.join(column_defs)},
                    exported_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            '''
            
            cursor.execute(create_table_sql)
            
            # Insert data
            placeholders = ', '.join(['%s' for _ in columns])
            insert_sql = f'INSERT INTO `{table_name}` ({", ".join([f"`{col}`" for col in columns])}) VALUES ({placeholders})'
            
            for item in self.data:
                values = [item.get(col) for col in columns]
                cursor.execute(insert_sql, values)
            
            conn.commit()
            cursor.close()
            conn.close()
            
            self.logger.info(f"Data exported to MySQL: {database}.{table_name}")
            return True
            
        except ImportError:
            self.logger.error("mysql-connector-python not installed. Run: pip install mysql-connector-python")
            return False
        except Exception as e:
            self.logger.error(f"MySQL export failed: {e}")
            return False
    
    def infer_sqlite_type(self, value: Any) -> str:
        """Infer SQLite data type"""
        if value is None:
            return 'TEXT'
        if isinstance(value, int):
            return 'INTEGER'
        if isinstance(value, float):
            return 'REAL'
        if isinstance(value, bool):
            return 'INTEGER'
        return 'TEXT'
    
    def infer_postgres_type(self, value: Any) -> str:
        """Infer PostgreSQL data type"""
        if value is None:
            return 'TEXT'
        if isinstance(value, int):
            return 'INTEGER'
        if isinstance(value, float):
            return 'FLOAT'
        if isinstance(value, bool):
            return 'BOOLEAN'
        if isinstance(value, dict) or isinstance(value, list):
            return 'JSONB'
        return 'TEXT'
    
    def infer_mysql_type(self, value: Any) -> str:
        """Infer MySQL data type"""
        if value is None:
            return 'TEXT'
        if isinstance(value, int):
            return 'INT'
        if isinstance(value, float):
            return 'FLOAT'
        if isinstance(value, bool):
            return 'BOOLEAN'
        if isinstance(value, dict) or isinstance(value, list):
            return 'JSON'
        
        # Check length for VARCHAR
        if isinstance(value, str):
            if len(value) > 255:
                return 'TEXT'
            return 'VARCHAR(255)'
        
        return 'TEXT'
    
    def load_from_json(self, json_file: str):
        """Load data from JSON file"""
        try:
            with open(json_file, 'r', encoding='utf-8') as f:
                self.data = json.load(f)
            self.metadata['total_records'] = len(self.data)
            self.logger.info(f"Loaded {len(self.data)} records from {json_file}")
            return True
        except Exception as e:
            self.logger.error(f"Failed to load JSON: {e}")
            return False
    
    def save_to_json(self, filename: str = None):
        """Save scraped data to JSON file"""
        if not filename:
            filename = f"scraped_data_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        
        output = {
            'metadata': self.metadata,
            'data': self.data
        }
        
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(output, f, indent=2, ensure_ascii=False)
        
        self.logger.info(f"Data saved to JSON: {filename}")
        return filename

def main():
    parser = argparse.ArgumentParser(description="Database Web Scraper - Export to Excel, MongoDB, or SQL")
    parser.add_argument("-u", "--url", help="Target URL to scrape")
    parser.add_argument("-m", "--mode", choices=["scrape", "table", "api", "load"], 
                       default="scrape", help="Scraping mode")
    parser.add_argument("-o", "--output", choices=["excel", "mongodb", "sqlite", "postgres", "mysql", "json"],
                       default="excel", help="Output format")
    parser.add_argument("-c", "--config", help="JSON configuration file")
    parser.add_argument("-t", "--table-index", type=int, default=0, help="Table index for table scraping")
    parser.add_argument("--api-params", help="API parameters as JSON string")
    
    # Database connection parameters
    parser.add_argument("--mongodb-uri", help="MongoDB connection URI")
    parser.add_argument("--mongodb-db", default="scraper_db", help="MongoDB database name")
    parser.add_argument("--mongodb-collection", default="scraped_data", help="MongoDB collection name")
    
    parser.add_argument("--sqlite-file", default="scraped_data.db", help="SQLite database file")
    parser.add_argument("--sqlite-table", default="scraped_data", help="SQLite table name")
    
    parser.add_argument("--postgres-host", default="localhost", help="PostgreSQL host")
    parser.add_argument("--postgres-db", default="scraper_db", help="PostgreSQL database")
    parser.add_argument("--postgres-user", help="PostgreSQL user")
    parser.add_argument("--postgres-password", help="PostgreSQL password")
    parser.add_argument("--postgres-table", default="scraped_data", help="PostgreSQL table name")
    
    parser.add_argument("--mysql-host", default="localhost", help="MySQL host")
    parser.add_argument("--mysql-db", default="scraper_db", help="MySQL database")
    parser.add_argument("--mysql-user", help="MySQL user")
    parser.add_argument("--mysql-password", help="MySQL password")
    parser.add_argument("--mysql-table", default="scraped_data", help="MySQL table name")
    
    args = parser.parse_args()
    
    scraper = DatabaseWebScraper()
    print(scraper.banner)
    
    # Configuration file support
    config = {}
    if args.config:
        try:
            with open(args.config, 'r') as f:
                config = json.load(f)
        except Exception as e:
            print(f"Error loading config: {e}")
            return
    
    # Perform scraping or loading
    if args.mode == "load":
        if not config.get('json_file'):
            print("Error: --mode load requires JSON file in config")
            return
        scraper.load_from_json(config['json_file'])
        
    elif args.mode == "table":
        if not args.url:
            print("Error: URL required for table scraping")
            return
        scraper.scrape_table(args.url, args.table_index)
        
    elif args.mode == "api":
        if not args.url:
            print("Error: API URL required")
            return
        params = json.loads(args.api_params) if args.api_params else None
        scraper.scrape_api_endpoint(args.url, params)
        
    else:  # scrape mode
        if not args.url or not config.get('selectors'):
            print("Error: URL and selectors configuration required for scraping mode")
            print("Example config.json:")
            print(json.dumps({
                "selectors": {
                    "_container": ".product",
                    "name": ".title",
                    "price": ".price",
                    "description": ".desc"
                },
                "pagination": {
                    "next_button": ".next-page",
                    "max_pages": 5,
                    "delay": 1
                }
            }, indent=2))
            return
        
        scraper.scrape_website(
            args.url,
            config['selectors'],
            config.get('pagination', {})
        )
    
    # Export data
    if args.output == "excel":
        scraper.export_to_excel()
        
    elif args.output == "mongodb":
        if not args.mongodb_uri:
            print("Error: MongoDB requires --mongodb-uri")
            return
        scraper.export_to_mongodb(
            args.mongodb_uri,
            args.mongodb_db,
            args.mongodb_collection
        )
        
    elif args.output == "sqlite":
        scraper.export_to_sqlite(args.sqlite_file, args.sqlite_table)
        
    elif args.output == "postgres":
        if not all([args.postgres_user, args.postgres_password]):
            print("Error: PostgreSQL requires --postgres-user and --postgres-password")
            return
        scraper.export_to_postgresql(
            args.postgres_host,
            args.postgres_db,
            args.postgres_user,
            args.postgres_password,
            args.postgres_table
        )
        
    elif args.output == "mysql":
        if not all([args.mysql_user, args.mysql_password]):
            print("Error: MySQL requires --mysql-user and --mysql-password")
            return
        scraper.export_to_mysql(
            args.mysql_host,
            args.mysql_db,
            args.mysql_user,
            args.mysql_password,
            args.mysql_table
        )
        
    elif args.output == "json":
        scraper.save_to_json()
    
    print(f"\n✅ Scraping complete! Total records: {len(scraper.data)}")

if __name__ == "__main__":
    main()
